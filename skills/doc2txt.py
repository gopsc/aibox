#!/usr/bin/env python3
"""
doc2txt - 将 Word 文档（.docx）转换为纯文本的工具
"""

import argparse
import sys
import os

from docx import Document


def extract_text(input_path, enable_tables=False, verbose=False):
    """
    从 docx 文件中提取文本内容。
    
    Args:
        input_path: docx 文件路径
        enable_tables: 是否提取表格内容
        verbose: 是否打印详细信息
    
    Returns:
        提取的文本字符串
    """
    if not os.path.exists(input_path):
        print(f"错误：文件不存在 - {input_path}", file=sys.stderr)
        sys.exit(1)
    
    if not input_path.endswith('.docx'):
        print(f"警告：文件不是 .docx 格式 - {input_path}", file=sys.stderr)
    
    doc = Document(input_path)
    output_lines = []
    
    para_count = 0
    table_count = 0
    
    # 提取段落文本
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            output_lines.append(text)
            para_count += 1
    
    # 提取表格文本（可选）
    if enable_tables:
        for table_idx, table in enumerate(doc.tables):
            table_count += 1
            if output_lines:
                output_lines.append("")  # 空行分隔
            output_lines.append(f"--- 表格 {table_idx + 1} ---")
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                output_lines.append(row_text)
    
    if verbose:
        print(f"[信息] 共处理 {len(doc.paragraphs)} 个段落", file=sys.stderr)
        print(f"[信息] 提取非空段落 {para_count} 个", file=sys.stderr)
        print(f"[信息] 共处理 {len(doc.tables)} 个表格", file=sys.stderr)
        if enable_tables:
            print(f"[信息] 提取表格 {table_count} 个", file=sys.stderr)
    
    return "\n".join(output_lines)


def main():
    parser = argparse.ArgumentParser(
        description="将 Word 文档（.docx）转换为纯文本",
        epilog="""
参数说明:
  -i, --input   输入的 .docx 文件路径（必填）
  -o, --output  输出 .txt 文件路径（可选，不指定则输出到 stdout）
  -t, --table   启用表格内容提取（默认仅提取段落文本）
  -v, --verbose 显示详细处理信息（段落数、表格数等）

使用示例:
  %(prog)s -i 报告.docx -o 报告.txt
  %(prog)s -i 报告.docx                        # 输出到终端
  %(prog)s -i 报告.docx -t                     # 同时提取表格
  %(prog)s -i 报告.docx -o 报告.txt -t -v      # 完整模式
        """,
        formatter_class=argparse.RawDescriptionHelpFormatter
    )
    
    parser.add_argument("-i", "--input", required=True, help="输入的 .docx 文件路径")
    parser.add_argument("-o", "--output", help="输出的 .txt 文件路径（可选，不指定则输出到 stdout）")
    parser.add_argument("-t", "--table", action="store_true", help="启用表格内容提取")
    parser.add_argument("-v", "--verbose", action="store_true", help="显示详细处理信息")
    
    args = parser.parse_args()
    
    text = extract_text(args.input, enable_tables=args.table, verbose=args.verbose)
    
    if args.output:
        with open(args.output, "w", encoding="utf-8") as f:
            f.write(text)
        if args.verbose:
            print(f"[信息] 已写入文件：{args.output}", file=sys.stderr)
    else:
        print(text)


if __name__ == "__main__":
    main()
